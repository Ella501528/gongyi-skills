# -*- coding: utf-8 -*-
"""志愿者名册与服务记录证明体检 —— 八条规则（逐条见 references/判定规则索引.md）

用法:
    python check_roster.py --input sample_data/sample_input.json --out ./名册体检报告.xlsx

输入: 一份 JSON（结构见 sample_data/sample_input.json，字段即填表说明）
输出: Excel 四表 —— 汇总 / 服务记录明细 / 异常清单 / 待补材料

未安装 openpyxl 时自动降级为 CSV（同目录，按表分文件，内容一致）。
命中的均为【待核项】，不是违规结论，须人工确认后再写入正式报告。

依据：《志愿服务记录与证明出具办法（试行）》民政部令第 67 号（第 3/4/5/6/7 条）、
      《慈善法》第 66 条、《志愿服务条例》第 19/20 条、四部门《关于规范志愿服务记录证明工作的指导意见》。
"""
import argparse
import csv
import datetime
import json
import os
import sys
from collections import Counter

SEV_ORDER = {"高": 0, "中": 1, "低": 2}

RULE_NAME = {
    "V1": "实名基本信息缺失",
    "V2": "服务记录七要素不全",
    "V3": "记录超期",
    "V4": "服务时长异常",
    "V5": "未录入指定信息系统",
    "V6": "证明应出未出或要素不全",
    "V7": "证明无编号或未盖章",
    "V8": "名册含未脱敏信息拟对外提供",
}
RULE_FIX = {
    "V1": "补收证件号等实名信息并核对；不得以内部编号替代",
    "V2": "当场补记七要素；补不出的写情况说明注明事后补记",
    "V3": "逐条补录说明；连续超期说明流程有断点，应改机制",
    "V4": "核对活动实际时长；负数与超上限须改账，不得虚记",
    "V5": "按统一数据标准补录民政部指定信息系统（待办，不阻断）",
    "V6": "核实后无偿、如实补出；要素不全的重出并作废原证",
    "V7": "补齐编号与公章；无公章者不具证明力",
    "V8": "先脱敏再对外；超出同意范围的信息不得提供",
}

REC_FIELDS = ["activity", "date", "location", "content", "hours", "organizer", "leader"]
PROOF_REQUIRED = ["hours", "content", "unit"]


def mask_name(name):
    if not name:
        return "（未填）"
    s = str(name).strip()
    return s if len(s) <= 1 else s[0] + "×" * (len(s) - 1)


def ensure_dir(path):
    d = os.path.dirname(os.path.abspath(path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)


def parse_date(s):
    try:
        return datetime.date.fromisoformat(str(s).strip())
    except Exception:
        return None


def business_days(d1, d2):
    """d1 → d2 之间的工作日数（不含 d1、含 d2）；d2 早于 d1 返回 0。"""
    if not d1 or not d2 or d2 < d1:
        return 0
    n, cur = 0, d1
    while cur < d2:
        cur += datetime.timedelta(days=1)
        if cur.weekday() < 5:
            n += 1
    return n


def run_rules(data, lag_days, max_hours):
    vols = data.get("volunteers", []) or []
    registered = set(data.get("system_registered") or [])
    roster_exported = bool(data.get("roster_exported"))
    default_registered = bool((data.get("org", {}) or {}).get("system_registered_default"))

    issues = []    # [志愿者号, 姓名, 规则, 严重度, 说明]
    detail = []    # 服务记录明细

    def add(vid, vname, rule, sev, note):
        issues.append([vid, vname, rule, sev, note])

    for v in vols:
        vid = v.get("id", "")
        vname = v.get("name", "")
        masked = mask_name(vname)

        # V1 实名基本信息缺失
        miss1 = [k for k in ["name", "id_no", "contact"] if not str(v.get(k, "")).strip()]
        if miss1:
            add(vid, masked, "V1", "高", "缺实名字段: " + "、".join(miss1))

        # 逐条服务记录
        for i, r in enumerate(v.get("records", []) or [], 1):
            omit = [k for k in REC_FIELDS if r.get(k) in (None, "")]
            if omit:
                add(vid, masked, "V2", "中", "第 %d 条服务记录缺: %s" % (i, "、".join(omit)))

            h = r.get("hours")
            if not isinstance(h, (int, float)) or h <= 0 or h > max_hours:
                add(vid, masked, "V4", "中", "第 %d 条服务时长异常: %s（上限 %s）" % (i, h, max_hours))

            d1 = parse_date(r.get("date"))
            d2 = parse_date(r.get("recorded_on"))
            if d1 and d2:
                bd = business_days(d1, d2)
                if bd > lag_days:
                    add(vid, masked, "V3", "中",
                        "第 %d 条：服务 %s → 记录 %s，超 %d 个工作日（阈值 %d）"
                        % (i, r.get("date"), r.get("recorded_on"), bd, lag_days))
            elif d1 is None or d2 is None:
                add(vid, masked, "V3", "低", "第 %d 条：缺服务日期或记录日期，无法判定是否超期" % i)

            detail.append([vid, masked, r.get("activity", ""), r.get("date", ""),
                           r.get("location", ""), r.get("content", ""), h,
                           r.get("organizer", ""), r.get("leader", ""),
                           r.get("recorded_on", "")])

        # V5 未录入指定信息系统
        if not (vid in registered or default_registered):
            add(vid, masked, "V5", "低", "未标记为已录入民政部指定志愿服务信息系统")

        # V6/V7 证明
        proofs = v.get("proofs", []) or []
        if v.get("proof_requested") and not proofs:
            add(vid, masked, "V6", "中", "已提出证明要求但无出具记录")
        for j, p in enumerate(proofs, 1):
            miss6 = [k for k in PROOF_REQUIRED if p.get(k) in (None, "")]
            if miss6:
                add(vid, masked, "V6", "中", "第 %d 份证明缺要素: %s" % (j, "、".join(miss6)))
            if not str(p.get("serial_no", "")).strip():
                add(vid, masked, "V7", "高", "第 %d 份证明无编号" % j)
            if p.get("sealed") is not True:
                add(vid, masked, "V7", "高", "第 %d 份证明未加盖公章（sealed≠true）" % j)

        # V8 名册含未脱敏信息拟对外提供
        if roster_exported:
            idno = str(v.get("id_no", ""))
            if idno and ("*" not in idno) and ("×" not in idno):
                add(vid, masked, "V8", "高", "名册拟对外提供，但证件号为全码未脱敏")

    return issues, detail


def build_tables(issues, detail, vols):
    high = sum(1 for i in issues if i[3] == "高")
    mid = sum(1 for i in issues if i[3] == "中")
    low = sum(1 for i in issues if i[3] == "低")

    summary = [
        ["项", "值"],
        ["志愿者人数", len(vols)],
        ["服务记录条数", len(detail)],
        ["命中总条数", len(issues)],
        ["高危条数（须人工确认）", high],
        ["中危条数", mid],
        ["低危条数（待办）", low],
        ["体检结论", "存在高危待核项，提交前须逐条确认" if high
         else ("仅中低危待核项，建议确认后放行" if (mid or low) else "未发现明显缺陷")],
        [],
        ["规则", "命中数"],
    ]
    cnt = Counter(i[2] for i in issues)
    for r in ["V1", "V2", "V3", "V4", "V5", "V6", "V7", "V8"]:
        summary.append(["%s %s" % (r, RULE_NAME[r]), cnt.get(r, 0)])

    sheet_detail = [["志愿者号", "姓名", "活动名称", "日期", "地点", "服务内容",
                     "服务时间(小时)", "活动组织单位", "活动负责人", "记录日期"]]
    sheet_detail.extend(detail)

    issues_sorted = sorted(issues, key=lambda x: SEV_ORDER.get(x[3], 9))
    sheet_issues = [["志愿者号", "姓名", "规则", "规则名", "严重度", "说明"]]
    for i in issues_sorted:
        sheet_issues.append([i[0], i[1], i[2], RULE_NAME[i[2]], i[3], i[4]])

    per_rule = Counter(i[2] for i in issues)
    fixes = [["规则", "需补材料／动作", "涉及条数"]]
    for r in ["V1", "V2", "V3", "V4", "V5", "V6", "V7", "V8"]:
        if per_rule.get(r):
            fixes.append([r + " " + RULE_NAME[r], RULE_FIX[r], per_rule[r]])

    return {"汇总": summary, "服务记录明细": sheet_detail, "异常清单": sheet_issues, "待补材料": fixes}


def write_xlsx(tables, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    first = True
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    warn_fill = PatternFill("solid", fgColor="FCE8E6")
    for name, rows in tables.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        if name == "异常清单":
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 5).value == "高":
                    for cc in range(1, 7):
                        ws.cell(r, cc).fill = warn_fill
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 52)
    wb.save(out_path)


def write_csv(tables, out_path):
    base, _ = os.path.splitext(out_path)
    written = []
    for name, rows in tables.items():
        p = "%s_%s.csv" % (base, name)
        ensure_dir(p)
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        written.append(p)
    return written


def main():
    ap = argparse.ArgumentParser(description="志愿者名册与服务记录证明体检")
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", default="./名册体检报告.xlsx")
    ap.add_argument("--lag-days", type=int, default=10,
                    help="记录超期阈值（工作日），默认 10")
    ap.add_argument("--max-hours", type=float, default=24,
                    help="单次服务时长上限（小时），默认 24")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print("输入文件不存在: %s" % args.input)
        sys.exit(2)

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    lag = args.lag_days
    if lag is None:
        lag = (data.get("org", {}) or {}).get("record_lag_workdays", 10)

    issues, detail = run_rules(data, lag, args.max_hours)
    tables = build_tables(issues, detail, data.get("volunteers", []) or [])
    ensure_dir(args.out)

    try:
        write_xlsx(tables, args.out)
        print("已生成 Excel: %s" % args.out)
    except ImportError:
        paths = write_csv(tables, args.out)
        print("未安装 openpyxl，已降级输出 CSV：")
        for p in paths:
            print("  " + p)

    print("志愿者 %d 人，服务记录 %d 条，命中 %d 条（高危 %d）" % (
        len(data.get("volunteers", []) or []), len(detail), len(issues),
        sum(1 for i in issues if i[3] == "高")))
    print("说明：命中项为待核项，不是违规结论，须人工确认后写入正式报告。")


if __name__ == "__main__":
    main()
