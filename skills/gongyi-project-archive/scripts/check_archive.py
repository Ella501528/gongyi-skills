# -*- coding: utf-8 -*-
"""项目档案归集体检 —— 五条规则（逐条见 references/判定规则索引.md）
并生成卷内目录（Word；无 python-docx 时降级为 HTML）。

用法:
    python check_archive.py --input sample_data/sample_input.json \
        --out ./归档体检报告.xlsx --docx ./卷内目录.docx

输入: 一份 JSON（结构见 sample_data/sample_input.json），列已收材料即可，
      支持按序号（如 "1.1"）或名称（如 "立项书"）填写。
输出: Excel 四表 —— 汇总 / 缺失清单 / 归集目录 / 补件顺序；--docx 另出卷内目录。

脚本只判【是否归集】，不判【要素是否合格】——要素完整性须人工复核。
"""
import argparse
import csv
import json
import os
import sys
from collections import Counter

NEED_ORDER = {"必要": 0, "尽量有": 1}

# 八类 35 项（来源：手册 3.3 节）
ARCHIVE = [
    # 一 立项类
    ("1.1", "立项类", "立项书（背景、目标、方案、预算、风险预案）", "必要", "第五十七条项目管理制度；协议成果要求", ["立项书", "项目立项书"]),
    ("1.2", "立项类", "立项审批记录／理事会决议", "必要", "第十二条治理结构与章程职权", ["立项审批", "理事会决议"]),
    ("1.3", "立项类", "宗旨与业务范围比对表", "必要", "第十一条、第十二条、第五十六条", ["业务范围比对表"]),
    ("1.4", "立项类", "资助标准与实施标准的论证记录", "必要", "第五十九条、第八十一条", ["资助标准论证", "实施标准论证"]),
    ("1.5", "立项类", "关联方识别与关联交易审议、公开记录", "必要", "第十四条", ["关联方识别记录", "关联交易记录"]),
    ("1.6", "立项类", "重大事项报备回执／报送记录", "尽量有", "属地重大事项报告要求", ["重大事项报备"]),
    # 二 协议类
    ("2.1", "协议类", "捐赠协议／资助协议（接收侧）", "必要", "第三十九条、第五十六条", ["捐赠协议", "资助协议"]),
    ("2.2", "协议类", "对外捐赠或资助协议（转赠侧）", "必要", "资助关系成立与用途约定", ["转赠协议"]),
    ("2.3", "协议类", "采购合同／服务合同", "必要", "采购制度与支出凭证要求", ["采购合同", "服务合同"]),
    ("2.4", "协议类", "合作方协议（权责、成果、宣传口径）", "必要", "第二十六条", ["合作协议", "合作方协议"]),
    ("2.5", "协议类", "用人／劳务协议", "尽量有", "劳动用工与劳务报酬规定", ["劳务协议", "用人协议"]),
    # 三 资金类
    ("3.1", "资金类", "拨付申请与审批单", "必要", "财务制度与审批权限", ["拨付申请", "付款审批单"]),
    ("3.2", "资金类", "银行流水／付款凭证", "必要", "第十二条、第五十四条", ["银行流水", "付款凭证"]),
    ("3.3", "资金类", "票据（捐赠票据、发票、收据）", "必要", "第三十八条", ["票据", "发票", "捐赠票据"]),
    ("3.4", "资金类", "全口径预算表、预算 vs 实际对比表", "必要", "第十三条；协议报告义务", ["预算表", "预算对比表"]),
    ("3.5", "资金类", "决算表（预算／实际／差异说明）", "必要", "第十三条、第七十八条", ["决算表"]),
    # 四 受益对象类
    ("4.1", "受益对象类", "资助标准与筛选条件（书面）", "必要", "第五十九条、第八十一条", ["筛选条件", "资助标准"]),
    ("4.2", "受益对象类", "受益对象名册", "必要", "资助协议与审计核查", ["名册", "受益人名册", "受助名单"]),
    ("4.3", "受益对象类", "资助金签收单／银行发放流水", "必要", "资金去向证明", ["签收单", "发放流水"]),
    ("4.4", "受益对象类", "变更或偏离说明与审批", "必要", "第五十六条", ["偏离说明", "变更说明"]),
    # 五 物资类
    ("5.1", "物资类", "采购比价记录（三方比价）", "必要", "采购制度与内控", ["比价记录", "三方比价"]),
    ("5.2", "物资类", "入库单", "必要", "第五十四条", ["入库单"]),
    ("5.3", "物资类", "出库发放单与签收单", "必要", "物资去向证明", ["出库单", "发放单"]),
    ("5.4", "物资类", "剩余物资处置审批与记录", "必要", "资产处置规定", ["剩余物资处置", "处置审批"]),
    # 六 过程类
    ("6.1", "过程类", "进度台账", "必要", "第五十七条", ["进度表", "项目台账"]),
    ("6.2", "过程类", "中期／阶段／年度汇报表", "必要", "协议报告义务；第十三条", ["中期汇报", "阶段汇报", "年度汇报", "项目月报"]),
    ("6.3", "过程类", "变更审批单（含用途变更备案或捐赠人同意）", "必要", "第五十六条", ["变更审批单"]),
    ("6.4", "过程类", "合作方定期对账记录", "尽量有", "内控与协议约定", ["对账记录"]),
    # 七 成效类
    ("7.1", "成效类", "成效数据（含前后对比）", "必要", "协议成果要求；评估举证", ["成效数据", "效果数据"]),
    ("7.2", "成效类", "受益对象反馈／回访记录", "必要", "效果最直接的表达之一", ["反馈记录", "回访记录"]),
    ("7.3", "成效类", "影像记录（含授权）", "尽量有", "成效佐证；个人信息保护规定", ["影像", "照片", "图片记录"]),
    ("7.4", "成效类", "第三方评估、媒体报道等外部印证", "尽量有", "效果的外印证", ["第三方评估", "媒体报道"]),
    # 八 结项类
    ("8.1", "结项类", "结项报告", "必要", "协议；评估与年报要求", ["结项报告"]),
    ("8.2", "结项类", "验收单／验收结论（含签署）", "必要", "协议约定的验收条件", ["验收单", "验收结论"]),
    ("8.3", "结项类", "档案归集目录（一项目一卷）", "尽量有", "档案管理制度", ["卷内目录", "归档目录"]),
]

CATS = ["立项类", "协议类", "资金类", "受益对象类", "物资类", "过程类", "成效类", "结项类"]

# 补件顺序（按被调取习惯）
FIX_ORDER = ["资金类", "受益对象类", "物资类", "立项类", "协议类", "过程类", "成效类", "结项类"]


def ensure_dir(path):
    d = os.path.dirname(os.path.abspath(path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)


def match(collected_raw):
    """把输入的已收材料匹配到清单项。返回 (命中序号集合, 未识别列表)。"""
    hits = set()
    unknown = []
    norm = [str(x).strip() for x in collected_raw if str(x).strip()]
    by_code = {a[0]: a for a in ARCHIVE}
    for token in norm:
        if token in by_code:
            hits.add(token)
            continue
        hit = None
        for code, cat, name, need, basis, alias in ARCHIVE:
            cands = [name] + list(alias)
            if any(token == c or (len(token) >= 3 and token in c) or (len(c) >= 3 and c in token)
                   for c in cands):
                hit = code
                break
        if hit:
            hits.add(hit)
        else:
            unknown.append(token)
    return hits, unknown


def run_rules(hits, unknown, extra_items):
    missing_need, missing_opt = [], []
    for code, cat, name, need, basis, alias in ARCHIVE:
        if code in hits:
            continue
        (missing_need if need == "必要" else missing_opt).append((code, cat, name, need, basis))

    issues = []
    for code, cat, name, need, basis in missing_need:
        issues.append([code, cat, name, "A1", "必要材料缺失", "高", "无此项即构成材料缺失"])
    for code, cat, name, need, basis in missing_opt:
        issues.append([code, cat, name, "A2", "建议材料缺失", "低", "有则举证更充分，一般不单独认定缺失"])

    # A3 整类空白
    for cat in CATS:
        codes = [a[0] for a in ARCHIVE if a[1] == cat]
        if not any(c in hits for c in codes):
            issues.append([codes[0], cat, "（整类）", "A3", "整类空白", "高", "该类 %d 项全部未收，重点核查" % len(codes)])

    # A4 结项前置不足
    jx_missing = [c for c in ["8.1", "8.2"] if c not in hits]
    if jx_missing:
        issues.append([",".join(jx_missing), "结项类", "（结项前置）", "A4", "结项前置不足", "高", "不得结项，先补结项报告与验收结论"])

    # A5 未识别材料
    for u in unknown:
        issues.append(["-", "（未识别）", u, "A5", "未识别材料", "中", "不在八类清单与别名表内，请人工确认"])

    extra = []
    for e in (extra_items or []):
        extra.append(["-", "（补充项）", str(e), "-", "清单外材料", "低", "已计入汇总，不改变八类主结构"])

    return issues, missing_need, missing_opt, extra


def build_tables(issues, missing_need, missing_opt, extra, hits, project):
    total = len(ARCHIVE)
    got = len(hits)
    summary = [
        ["项", "值"],
        ["项目名称", project.get("name", "（未填）")],
        ["项目期间", project.get("period", "（未填）")],
        ["清单项目数", total],
        ["已归集", got],
        ["缺失（必要）", len(missing_need)],
        ["缺失（尽量有）", len(missing_opt)],
        ["清单外补充项", len(extra)],
        ["归集率", "%.0f%%" % (got / total * 100)],
        ["结项前置", "齐备" if not [i for i in issues if i[3] == "A4"] else "不足，不得结项"],
        [],
        ["类", "已收", "必要缺", "尽量有缺"],
    ]
    for cat in CATS:
        c_all = [a for a in ARCHIVE if a[1] == cat]
        c_got = sum(1 for a in c_all if a[0] in hits)
        summary.append([cat, "%d/%d" % (c_got, len(c_all)),
                        sum(1 for a in missing_need if a[1] == cat),
                        sum(1 for a in missing_opt if a[1] == cat)])

    sheet_missing = [["序号", "类", "材料", "必要性", "要求或依据"]]
    for row in missing_need + missing_opt:
        sheet_missing.append(list(row))

    sheet_dir = [["序号", "类", "材料", "必要性", "状态", "页码"]]
    for code, cat, name, need, basis, alias in ARCHIVE:
        sheet_dir.append([code, cat, name, need, "已归集" if code in hits else "缺失", ""])
    for row in extra:
        sheet_dir.append([row[0], row[1], row[2], "", row[4], ""])

    per_cat_need = Counter(a[1] for a in missing_need)
    per_cat_opt = Counter(a[1] for a in missing_opt)
    sheet_fix = [["顺序", "类", "必要缺", "尽量有缺", "先补什么"]]
    tips = {
        "资金类": "拨付审批、流水、票据、预算与决算",
        "受益对象类": "筛选条件、名册、签收或流水、变更说明",
        "物资类": "比价记录、入库单、出库签收、剩余处置",
        "立项类": "立项书、审批决议、业务范围比对、标准论证、关联方记录",
        "协议类": "接收侧与转赠侧协议、采购与合作协议",
        "过程类": "进度台账、中期汇报、变更审批、对账记录",
        "成效类": "成效数据、反馈回访、影像与外印证",
        "结项类": "结项报告、验收结论、卷内目录",
    }
    for i, cat in enumerate(FIX_ORDER, 1):
        sheet_fix.append([i, cat, per_cat_need.get(cat, 0), per_cat_opt.get(cat, 0), tips.get(cat, "")])

    high = sum(1 for i in issues if i[5] == "高")
    summary.append([])
    summary.append(["规则", "命中数"])
    cnt = Counter(i[3] for i in issues)
    for r, rn in [("A1", "必要材料缺失"), ("A2", "建议材料缺失"), ("A3", "整类空白"),
                  ("A4", "结项前置不足"), ("A5", "未识别材料")]:
        summary.append(["%s %s" % (r, rn), cnt.get(r, 0)])
    summary.append([])
    summary.append(["体检结论", "存在高危待核项，结项或迎检前须逐条补正" if high
                    else ("仅低危提示项" if issues else "未发现缺口")])

    return {"汇总": summary, "缺失清单": sheet_missing, "归集目录": sheet_dir, "补件顺序": sheet_fix}


def write_xlsx(tables, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    first = True
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    warn_fill = PatternFill("solid", fgColor="FCE8E6")
    ok_fill = PatternFill("solid", fgColor="E6F4EA")
    for name, rows in tables.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        if name == "归集目录":
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, 5).value
                fill = ok_fill if v == "已归集" else (warn_fill if v == "缺失" else None)
                if fill:
                    ws.cell(r, 5).fill = fill
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 56)
    wb.save(out_path)


def write_csv(tables, out_path):
    base, _ = os.path.splitext(out_path)
    written = []
    for name, rows in tables.items():
        p = "%s_%s.csv" % (base, name)
        ensure_dir(p)
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        written.append(p)
    return written


def write_index_docx(rows, out_path, project):
    """rows: [序号, 类, 材料, 必要性, 状态, 页码]"""
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(10.5)
    h = doc.add_heading("项目档案卷内目录", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("项目名称：%s　　项目期间：%s" % (
        project.get("name", "＿＿＿＿"), project.get("period", "＿＿＿＿")))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2 = doc.add_paragraph("载体：纸质／电子　　卷号：＿＿＿　　编制人：＿＿＿　　编制日期：＿＿＿＿")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    tb = doc.add_table(rows=1, cols=6)
    tb.style = "Table Grid"
    hdr = ["序号", "类", "材料名称", "必要性", "状态", "页码"]
    for i, t in enumerate(hdr):
        tb.rows[0].cells[i].text = t
    for row in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    doc.add_paragraph("")
    doc.add_paragraph("说明：状态为「已归集／缺失」；页码栏由装订后手工填写。"
                      "缺项须按补件顺序补齐后方可结项。")
    ensure_dir(out_path)
    doc.save(out_path)


def write_index_html(rows, out_path, project):
    trs = "".join("<tr>" + "".join("<td>%s</td>" % ("" if c is None else c) for c in r) + "</tr>" for r in rows)
    html = """<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">
<title>项目档案卷内目录</title><style>
body{font-family:"Microsoft YaHei",system-ui,sans-serif;margin:32px;color:#1c1c1e}
h1{font-size:20px;text-align:center}
p.sub{text-align:center;color:#555;font-size:13px}
table{border-collapse:collapse;width:100%;margin-top:16px;font-size:13px}
th,td{border:1px solid #bbb;padding:7px 9px;text-align:left}
th{background:#eef3fa}
td:nth-child(1),td:nth-child(4),td:nth-child(5),td:nth-child(6){white-space:nowrap}
</style></head><body>
<h1>项目档案卷内目录</h1>
<p class="sub">项目名称：%s　｜　项目期间：%s</p>
<p class="sub">载体：纸质／电子　｜　卷号：＿＿＿　｜　编制人：＿＿＿　｜　编制日期：＿＿＿＿</p>
<table><thead><tr><th>序号</th><th>类</th><th>材料名称</th><th>必要性</th><th>状态</th><th>页码</th></tr></thead>
<tbody>%s</tbody></table>
<p class="sub">说明：状态为「已归集／缺失」；页码栏由装订后手工填写。缺项须按补件顺序补齐后方可结项。</p>
</body></html>""" % (project.get("name", "＿＿＿＿"), project.get("period", "＿＿＿＿"), trs)
    ensure_dir(out_path)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)


def main():
    ap = argparse.ArgumentParser(description="项目档案归集体检 + 卷内目录生成")
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", default="./归档体检报告.xlsx")
    ap.add_argument("--docx", default=None, help="卷内目录输出路径（.docx）")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print("输入文件不存在: %s" % args.input)
        sys.exit(2)

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    project = data.get("project", {}) or {}
    hits, unknown = match(data.get("collected", []) or [])
    issues, mn, mo, extra = run_rules(hits, unknown, data.get("extra_items"))
    tables = build_tables(issues, mn, mo, extra, hits, project)
    ensure_dir(args.out)

    try:
        write_xlsx(tables, args.out)
        print("已生成 Excel: %s" % args.out)
    except ImportError:
        for p in write_csv(tables, args.out):
            print("  " + p)
        print("未安装 openpyxl，已降级输出 CSV。")

    if args.docx:
        try:
            write_index_docx(tables["归集目录"][1:], args.docx, project)
            print("已生成 Word 卷内目录: %s" % args.docx)
        except ImportError:
            html_path = os.path.splitext(args.docx)[0] + ".html"
            write_index_html(tables["归集目录"][1:], html_path, project)
            print("未安装 python-docx，已降级输出 HTML 卷内目录: %s" % html_path)

    print("已归集 %d/%d，必要缺 %d，尽量有缺 %d，未识别 %d" % (
        len(hits), len(ARCHIVE), len(mn), len(mo), len(unknown)))
    print("说明：脚本只判是否归集，要素完整性须人工复核。")


if __name__ == "__main__":
    main()
