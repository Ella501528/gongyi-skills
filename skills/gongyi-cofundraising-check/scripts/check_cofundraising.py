# -*- coding: utf-8 -*-
"""合作募捐合规与对账体检 —— 八条规则（逐条见 references/判定规则索引.md）

用法:
    python check_cofundraising.py --input sample_data/sample_input.json --out ./合作募捐体检报告.xlsx

输入: 一份 JSON，合作募捐项目清单 cofundraisings（结构见 sample_data/sample_input.json）
输出: Excel 四表 —— 汇总 / 募捐明细 / 异常清单 / 待补材料

未安装 openpyxl 时自动降级为 CSV（同目录，按表分文件，内容一致）。
命中的均为【待核项】，不是违规结论，须人工确认。
本技能只管公开募捐【合作】的专门合规；合作方一般准入见 gongyi-partner-onboarding。

依据：《慈善法》第二十四条、第二十五条、第二十六条、第五十三条、第五十四条、第五十六条。
"""
import argparse
import csv
import json
import os
import sys
from collections import Counter

SEV_ORDER = {"高": 0, "中": 1, "低": 2}

RULE_NAME = {
    "C1": "公开募捐资格缺位或未依托",
    "C2": "未签书面协议",
    "C3": "协议未载明关键事项",
    "C4": "未对合作方进行评估",
    "C5": "未履行监督义务",
    "C6": "募得款物未全部纳入本机构账户",
    "C7": "募捐方案未按规定公开",
    "C8": "支出超出募捐方案与协议范围",
}
RULE_FIX = {
    "C1": "调整为不具资格方与具资格慈善组织合作，款物由具资格方管理；否则不得公开募捐",
    "C2": "补签书面协议",
    "C3": "补齐募捐目的、募捐方案、款物用途、各方权利义务",
    "C4": "补合作方评估并留存依据",
    "C5": "建立阶段进展、款物使用与对账留痕",
    "C6": "募得款物全部纳入具资格慈善组织账户，单独核算；已体外收取的限期归集并说明",
    "C7": "按规定备案并公开；未公开不得对外开募",
    "C8": "停止超范围使用；需变更用途的报原备案民政部门备案或征得捐赠人同意",
}
AGREEMENT_REQUIRED = ["募捐目的", "募捐方案", "款物用途", "各方权利义务"]


def mask_org(name):
    if not name:
        return "（未填）"
    s = str(name).strip()
    return s if len(s) <= 2 else s[0] + "×" * (len(s) - 1)


def ensure_dir(path):
    d = os.path.dirname(os.path.abspath(path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)


def run_rules(data):
    org = data.get("org", {}) or {}
    org_has = bool(org.get("has_public_fundraising_qualification"))
    items = data.get("cofundraisings", []) or []

    issues = []   # [项目号, 名称, 规则, 严重度, 说明]
    detail = []

    def add(fid, fname, rule, sev, note):
        issues.append([fid, fname, rule, sev, note])

    for it in items:
        fid = it.get("id", "")
        fname = mask_org(it.get("name", ""))
        partner_has = bool(it.get("partner_has_qualification"))
        relying = bool(it.get("partner_relying_on_qualified"))

        # C1 资格前提
        if not org_has and not partner_has:
            add(fid, fname, "C1", "高", "合作双方均无公开募捐资格，不得开展公开募捐")
        elif (not partner_has) and (not relying):
            add(fid, fname, "C1", "高", "合作方不具资格但未依托具资格的慈善组织")

        # C2/C3 协议
        if not it.get("agreement_signed", False):
            add(fid, fname, "C2", "高", "未签书面协议")
        else:
            miss = [f for f in AGREEMENT_REQUIRED if f not in (it.get("agreement_fields") or [])]
            if miss:
                add(fid, fname, "C3", "中", "协议未载明: " + "、".join(miss))

        # C4 评估
        if not it.get("partner_evaluated", False):
            add(fid, fname, "C4", "中", "未对合作方进行评估")

        # C5 监督
        if not it.get("supervised", False):
            add(fid, fname, "C5", "中", "未见监督留痕")

        # C6 入账
        if not it.get("funds_to_org_account", False):
            add(fid, fname, "C6", "高", "募得款物未全部纳入本机构账户")

        # C7 方案公开
        if not it.get("scheme_disclosed", False):
            add(fid, fname, "C7", "中", "募捐方案未按规定公开")

        # C8 支出范围
        if not it.get("expenses_within_scheme", False):
            add(fid, fname, "C8", "高", "支出超出募捐方案与协议范围")

        detail.append([fid, fname,
                       "是" if partner_has else "否",
                       "是" if relying else "否",
                       "是" if it.get("agreement_signed") else "否",
                       "是" if it.get("partner_evaluated") else "否",
                       "是" if it.get("supervised") else "否",
                       "是" if it.get("funds_to_org_account") else "否",
                       "是" if it.get("scheme_disclosed") else "否",
                       "是" if it.get("expenses_within_scheme") else "否"])

    return issues, detail


def build_tables(issues, detail, items):
    high = sum(1 for i in issues if i[3] == "高")
    mid = sum(1 for i in issues if i[3] == "中")
    low = sum(1 for i in issues if i[3] == "低")

    summary = [
        ["项", "值"],
        ["合作募捐项目数", len(items)],
        ["命中总条数", len(issues)],
        ["高危条数（须人工确认）", high],
        ["中危条数", mid],
        ["低危条数（待办）", low],
        ["体检结论", "存在高危待核项，开募前须逐条确认" if high
         else ("仅中低危待核项，建议确认后放行" if (mid or low) else "未发现明显缺陷")],
        [],
        ["规则", "命中数"],
    ]
    cnt = Counter(i[2] for i in issues)
    for r in ["C1", "C2", "C3", "C4", "C5", "C6", "C7", "C8"]:
        summary.append(["%s %s" % (r, RULE_NAME[r]), cnt.get(r, 0)])

    sheet_detail = [["项目号", "项目名称", "合作方具资格", "已依托具资格方", "已签协议",
                     "已评估", "已监督", "款物入账", "方案已公开", "支出在方案内"]]
    sheet_detail.extend(detail)

    issues_sorted = sorted(issues, key=lambda x: SEV_ORDER.get(x[3], 9))
    sheet_issues = [["项目号", "项目名称", "规则", "规则名", "严重度", "说明"]]
    for i in issues_sorted:
        sheet_issues.append([i[0], i[1], i[2], RULE_NAME[i[2]], i[3], i[4]])

    per_rule = Counter(i[2] for i in issues)
    fixes = [["规则", "需补材料／动作", "涉及条数"]]
    for r in ["C1", "C2", "C3", "C4", "C5", "C6", "C7", "C8"]:
        if per_rule.get(r):
            fixes.append([r + " " + RULE_NAME[r], RULE_FIX[r], per_rule[r]])

    return {"汇总": summary, "募捐明细": sheet_detail, "异常清单": sheet_issues, "待补材料": fixes}


def write_xlsx(tables, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    first = True
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    warn_fill = PatternFill("solid", fgColor="FCE8E6")
    for name, rows in tables.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        if name == "异常清单":
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 5).value == "高":
                    for cc in range(1, 7):
                        ws.cell(r, cc).fill = warn_fill
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 52)
    wb.save(out_path)


def write_csv(tables, out_path):
    base, _ = os.path.splitext(out_path)
    written = []
    for name, rows in tables.items():
        p = "%s_%s.csv" % (base, name)
        ensure_dir(p)
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        written.append(p)
    return written


def main():
    ap = argparse.ArgumentParser(description="合作募捐合规与对账体检")
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", default="./合作募捐体检报告.xlsx")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print("输入文件不存在: %s" % args.input)
        sys.exit(2)

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    issues, detail = run_rules(data)
    tables = build_tables(issues, detail, data.get("cofundraisings", []) or [])
    ensure_dir(args.out)

    try:
        write_xlsx(tables, args.out)
        print("已生成 Excel: %s" % args.out)
    except ImportError:
        paths = write_csv(tables, args.out)
        print("未安装 openpyxl，已降级输出 CSV：")
        for p in paths:
            print("  " + p)

    print("合作募捐项目 %d 个，命中 %d 条（高危 %d）" % (
        len(data.get("cofundraisings", []) or []), len(issues),
        sum(1 for i in issues if i[3] == "高")))
    print("说明：命中项为待核项，不是违规结论，须人工确认后使用。")


if __name__ == "__main__":
    main()
