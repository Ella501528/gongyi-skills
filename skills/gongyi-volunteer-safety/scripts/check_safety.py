# -*- coding: utf-8 -*-
"""志愿者风险与保险体检 —— 八条规则（逐条见 references/判定规则索引.md）

用法:
    python check_safety.py --input sample_data/sample_input.json --out ./保障体检报告.xlsx --min-coverage 100000

输入: 一份 JSON，含 activities（活动与投保）、participants（参加人员）、incidents（出险记录）
输出: Excel 四表 —— 汇总 / 保障明细 / 异常清单 / 待补材料

未安装 openpyxl 时自动降级为 CSV（同目录，按表分文件，内容一致）。
命中的均为【待核项】，不是违规结论，须人工确认。
本技能不推荐具体保险产品，也不替代保险专业人士的核保意见。

依据：《慈善法》第 68/119 条、《志愿服务条例》第 17 条、第三十三条（第二类情形）。
"""
import argparse
import csv
import datetime
import json
import os
import sys
from collections import Counter

SEV_ORDER = {"高": 0, "中": 1, "低": 2}

RULE_NAME = {
    "S1": "可能发生人身危险的活动未投保",
    "S2": "投保时点晚于服务开始",
    "S3": "保额低于机构最低标准",
    "S4": "参加人未在保单名单内",
    "S5": "无投保凭证留档",
    "S6": "大型活动主办方未按约定投保",
    "S7": "出险未按流程处置",
    "S8": "要求志愿者自担投保费用",
}
RULE_FIX = {
    "S1": "服务开始前完成投保；未投保不得安排上岗",
    "S2": "投保须在服务开始前完成；已发生者书面说明并整改",
    "S3": "按机构标准调整保额；无标准的先定标准",
    "S4": "补入保单名单或调整排班；漏人是最高发漏洞",
    "S5": "留存保单号、承保公司、保额、生效日、被保险人名单",
    "S6": "协议写清由谁投保与保额；主办方未履行的要求其履行或组织先行投保并追偿",
    "S7": "补救助记录、事故经过、索赔材料；因第三人原因的协助索赔",
    "S8": "组织承担应负部分；志愿者自愿自费加保不替代组织义务",
}


def mask_name(name):
    if not name:
        return "（未填）"
    s = str(name).strip()
    return s if len(s) <= 1 else s[0] + "×" * (len(s) - 1)


def ensure_dir(path):
    d = os.path.dirname(os.path.abspath(path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)


def parse_date(s):
    try:
        return datetime.date.fromisoformat(str(s).strip())
    except Exception:
        return None


def run_rules(data, min_coverage):
    acts = {a.get("id"): a for a in (data.get("activities", []) or [])}
    parts = data.get("participants", []) or []
    incidents = data.get("incidents", []) or []

    issues = []   # [活动, 对象, 规则, 严重度, 说明]
    detail = []
    part_by_act = {}
    for p in parts:
        part_by_act.setdefault(p.get("activity_id"), []).append(p)

    def add(act, who, rule, sev, note):
        issues.append([act, who, rule, sev, note])

    for a in acts.values():
        aid = a.get("id", "")
        aname = a.get("name", aid)
        risk = a.get("risk_level", "低")
        dangerous = bool(a.get("dangerous"))
        ins = a.get("insurance", {}) or {}
        insured = bool(ins.get("insured"))
        coverage = ins.get("coverage")
        insured_on = parse_date(ins.get("insured_on"))
        start = parse_date(a.get("start_date"))
        listed = set(ins.get("insured_volunteer_ids") or [])

        # S1
        if dangerous and not insured:
            add(aname, "（活动级）", "S1", "高", "可能发生人身危险的活动未投保")

        if insured:
            # S2 时点
            if insured_on and start and insured_on > start:
                add(aname, "（活动级）", "S2", "高",
                    "保单生效 %s 晚于活动开始 %s，未在服务前完成" % (ins.get("insured_on"), a.get("start_date")))
            # S3 保额
            if min_coverage is None:
                add(aname, "（活动级）", "S3", "低", "未设定最低保额标准，无法判定")
            elif isinstance(coverage, (int, float)) and coverage < min_coverage:
                add(aname, "（活动级）", "S3", "中",
                    "保额 %s 低于最低标准 %s" % (coverage, min_coverage))
            # S5 凭证
            if not ins.get("voucher_kept", False):
                add(aname, "（活动级）", "S5", "中", "已投保但无凭证留档")

        # S4 名单
        if insured:
            for p in part_by_act.get(aid, []):
                if p.get("volunteer_id") not in listed:
                    add(aname, mask_name(p.get("volunteer_name")), "S4", "高",
                        "参加人 %s 不在保单名单内" % p.get("volunteer_id"))

        # S6 大型活动主办方投保
        if a.get("is_large_event") and not a.get("organizer_insured", False):
            add(aname, "（主办方）", "S6", "高", "大型公益活动主办方未按约定投保")

        # S8 费用转嫁
        if a.get("volunteer_bears_insurance_cost", False):
            add(aname, "（活动级）", "S8", "中", "要求志愿者自担投保费用")

        detail.append([aid, aname, risk, "是" if dangerous else "否",
                       "是" if insured else "否", ins.get("insured_on", ""),
                       coverage, len(listed), len(part_by_act.get(aid, [])),
                       "是" if a.get("organizer_insured") else "否"])

    # S7 出险处置
    for it in incidents:
        aid = it.get("activity_id", "")
        aname = (acts.get(aid) or {}).get("name", aid)
        who = mask_name(it.get("volunteer_name"))
        missing = []
        if not it.get("rescued", False):
            missing.append("未及时救助")
        if not it.get("record_kept", False):
            missing.append("未留事故记录")
        if not it.get("claim_assisted", False):
            missing.append("未协助索赔")
        if missing:
            add(aname, who, "S7", "高", "；".join(missing))

    return issues, detail


def build_tables(issues, detail, acts):
    high = sum(1 for i in issues if i[3] == "高")
    mid = sum(1 for i in issues if i[3] == "中")
    low = sum(1 for i in issues if i[3] == "低")

    summary = [
        ["项", "值"],
        ["活动数", len(acts)],
        ["命中总条数", len(issues)],
        ["高危条数（须人工确认）", high],
        ["中危条数", mid],
        ["低危条数（待办）", low],
        ["体检结论", "存在高危待核项，上岗前须逐条确认" if high
         else ("仅中低危待核项，建议确认后放行" if (mid or low) else "未发现明显缺陷")],
        [],
        ["规则", "命中数"],
    ]
    cnt = Counter(i[2] for i in issues)
    for r in ["S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8"]:
        summary.append(["%s %s" % (r, RULE_NAME[r]), cnt.get(r, 0)])

    sheet_detail = [["活动号", "活动", "风险级别", "可能人身危险", "已投保", "生效日",
                     "保额", "保单名单人数", "参加人数", "主办方已投保"]]
    sheet_detail.extend(detail)

    issues_sorted = sorted(issues, key=lambda x: SEV_ORDER.get(x[3], 9))
    sheet_issues = [["活动", "对象", "规则", "规则名", "严重度", "说明"]]
    for i in issues_sorted:
        sheet_issues.append([i[0], i[1], i[2], RULE_NAME[i[2]], i[3], i[4]])

    per_rule = Counter(i[2] for i in issues)
    fixes = [["规则", "需补材料／动作", "涉及条数"]]
    for r in ["S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8"]:
        if per_rule.get(r):
            fixes.append([r + " " + RULE_NAME[r], RULE_FIX[r], per_rule[r]])

    return {"汇总": summary, "保障明细": sheet_detail, "异常清单": sheet_issues, "待补材料": fixes}


def write_xlsx(tables, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    first = True
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    warn_fill = PatternFill("solid", fgColor="FCE8E6")
    for name, rows in tables.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        if name == "异常清单":
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 5).value == "高":
                    for cc in range(1, 7):
                        ws.cell(r, cc).fill = warn_fill
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 52)
    wb.save(out_path)


def write_csv(tables, out_path):
    base, _ = os.path.splitext(out_path)
    written = []
    for name, rows in tables.items():
        p = "%s_%s.csv" % (base, name)
        ensure_dir(p)
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        written.append(p)
    return written


def main():
    ap = argparse.ArgumentParser(description="志愿者风险与保险体检")
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", default="./保障体检报告.xlsx")
    ap.add_argument("--min-coverage", type=float, default=None,
                    help="机构最低保额标准；不传则读输入 org.min_coverage，仍无则报无法判定")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print("输入文件不存在: %s" % args.input)
        sys.exit(2)

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    mc = args.min_coverage
    if mc is None:
        mc = (data.get("org", {}) or {}).get("min_coverage")
        if mc == "":
            mc = None

    issues, detail = run_rules(data, mc)
    tables = build_tables(issues, detail, {a.get("id"): a for a in (data.get("activities", []) or [])})
    ensure_dir(args.out)

    try:
        write_xlsx(tables, args.out)
        print("已生成 Excel: %s" % args.out)
    except ImportError:
        paths = write_csv(tables, args.out)
        print("未安装 openpyxl，已降级输出 CSV：")
        for p in paths:
            print("  " + p)

    print("活动 %d 个，命中 %d 条（高危 %d）" % (
        len(data.get("activities", []) or []), len(issues),
        sum(1 for i in issues if i[3] == "高")))
    print("说明：命中项为待核项，不是违规结论，须人工确认后使用；本技能不推荐具体保险产品。")


if __name__ == "__main__":
    main()
