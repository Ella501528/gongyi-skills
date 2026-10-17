# -*- coding: utf-8 -*-
"""受益对象名册体检 —— 八条规则（逐条见 references/判定规则索引.md）

用法:
    python check_roster.py --input sample_data/sample_input.json --out ./体检报告.xlsx

输入: 一份 JSON（结构见 sample_data/sample_input.json，字段即填表说明）
输出: Excel 四表 —— 汇总 / 缺陷清单 / 偏离提示 / 待补材料

未安装 openpyxl 时自动降级为 CSV（同目录，按表分文件，内容一致）。
命中的均为【待核项】，不是违规结论，须人工确认后再写入正式报告。
"""
import argparse
import csv
import json
import os
import sys
from collections import Counter

# ---------------------------------------------------------------- 常量

SEV_ORDER = {"高": 0, "中": 1, "低": 2}

RULE_NAME = {
    "R1": "追溯字段缺失",
    "R2": "资格依据缺失",
    "R3": "最小必要越界",
    "R4": "未成年人未尽从严",
    "R5": "签收缺失",
    "R6": "给付凭证不匹配",
    "R7": "人群/区域分布偏离",
    "R8": "变更无审批",
}
RULE_FIX = {
    "R1": "补齐姓名、发放方式、区域（到乡镇/社区即可）",
    "R2": "补资格类别与认定依据（审核表或评审记录）",
    "R3": "删除或脱敏超范围字段；保留的须写明用途与保管方式",
    "R4": "补监护人同意记录；未补前不得发放",
    "R5": "当日补签；补不出的写情况说明",
    "R6": "按给付形态补对应凭证（现金看流水/签收，实物看签收）",
    "R7": "人工确认是否属客观原因；属执行的履行变更审批",
    "R8": "补原因记录与变更审批；在报告如实反映",
}


def mask_name(name):
    """只保留姓氏，其余打码，避免报告外泄完整身份信息。"""
    if not name:
        return "（未填）"
    name = str(name).strip()
    if len(name) <= 1:
        return name
    return name[0] + "×" * (len(name) - 1)


# ---------------------------------------------------------------- 规则引擎

def run_rules(data, thr_gender, thr_age):
    project = data.get("project", {})
    target = project.get("target_group", {}) or {}
    people = data.get("beneficiaries", []) or []
    extended = set(data.get("extended_fields", []) or [])
    declared_purpose = set((project.get("fields_purpose") or {}).keys())

    issues = []          # (编号, 掩码姓名, 规则, 严重度, 说明)
    deviations = []      # (维度, 目标, 实际, 提示)

    def add(pid, pname, rule, sev, note):
        issues.append([pid, mask_name(pname), rule, sev, note])

    for p in people:
        pid = p.get("id", "")
        pname = p.get("name", "")

        # R1 追溯字段缺失
        miss = [k for k in ("name", "payout", "region") if not p.get(k)]
        if miss:
            add(pid, pname, "R1", "高", "缺: " + "、".join(miss))

        # R2 资格依据缺失
        if not p.get("qualification") or not p.get("qualification_basis"):
            add(pid, pname, "R2", "高", "资格类别或认定依据未填")

        # R4 未成年人未尽从严
        if p.get("minor") is True and p.get("guardian_consent") is not True:
            add(pid, pname, "R4", "高", "未成年人但无监护人同意记录")

        # R5 签收缺失 / R6 凭证不匹配
        payout = p.get("payout", "")
        ev = p.get("payout_evidence", "")
        signed = p.get("signed")
        has_proof = (ev in ("bank", "signoff")) or (signed is True)
        if not has_proof:
            add(pid, pname, "R5", "高", "无签收、也无流水/签到等等价凭证")
        if payout == "现金" and ev not in ("bank", "signoff", "流水"):
            add(pid, pname, "R6", "高", "现金发放缺银行流水或签收单")
        if payout in ("实物", "物资") and signed is not True:
            add(pid, pname, "R6", "高", "实物发放缺签收单")

        # R8 变更无审批
        if p.get("status") in ("退出", "替换") and p.get("change_approved") is not True:
            add(pid, pname, "R8", "中", "状态为%s但无变更审批" % p.get("status"))

    # R3 最小必要越界（整体层）
    if extended:
        add("-", "-", "R3", "中", "采集了超范围字段: " + "、".join(sorted(extended)))
    undeclared = []
    for p in people:
        for k, v in p.items():
            if v in (None, "", False):
                continue
            if k not in declared_purpose and k not in (
                "id", "name", "sex", "age", "region", "qualification",
                "qualification_basis", "payout", "payout_amount",
                "payout_evidence", "signed", "minor", "guardian_consent",
                "status", "change_approved",
            ):
                undeclared.append(k)
    if undeclared:
        add("-", "-", "R3", "中", "存在未说明用途的字段: " + "、".join(sorted(set(undeclared))))

    # R7 分布偏离
    total = len(people)
    if total == 0:
        deviations.append(["人数", "-", "0", "名册为空，无法评估"])
        return issues, deviations, total

    tgt_sex = target.get("sex")
    if tgt_sex:
        cnt = Counter(p.get("sex") for p in people)
        ratio = cnt.get(tgt_sex, 0) / total
        deviations.append([
            "性别构成", "目标性别 %s" % tgt_sex,
            "占比 %.0f%%（实到 %d/%d）" % (ratio * 100, cnt.get(tgt_sex, 0), total),
            "低于阈值 %.0f%% —— 提示人群可能偏离" % (thr_gender * 100) if ratio < thr_gender
            else "在阈值内",
        ])

    amin, amax = target.get("age_min"), target.get("age_max")
    if amin is not None and amax is not None:
        ages = [p.get("age") for p in people if isinstance(p.get("age"), (int, float))]
        if ages:
            out = [a for a in ages if a < amin or a > amax]
            ratio = len(out) / total
            deviations.append([
                "年龄区间", "目标 %s–%s 岁" % (amin, amax),
                "超区间 %d 人（%.0f%%）" % (len(out), ratio * 100),
                "高于阈值 %.0f%% —— 提示标准可能偏离" % (thr_age * 100) if ratio > thr_age
                else "在阈值内",
            ])

    regions = target.get("regions")
    if regions:
        outside = [p for p in people if p.get("region") and p.get("region") not in regions]
        deviations.append([
            "区域", "目标区域 " + "、".join(regions),
            "区域外 %d 人" % len(outside),
            "存在目标区域外对象，逐条人工确认" if outside else "无区域外对象",
        ])

    return issues, deviations, total


# ---------------------------------------------------------------- 输出

def build_tables(issues, deviations, total, project):
    high = sum(1 for i in issues if i[3] == "高")
    mid = sum(1 for i in issues if i[3] == "中")
    summary = [
        ["项", "值"],
        ["项目名称", project.get("name", "（未填）")],
        ["名册人数", total],
        ["命中总条数", len(issues)],
        ["高危条数（须人工确认）", high],
        ["中危条数", mid],
        ["体检结论", "存在高危待核项，发放或提交前须逐条确认" if high
         else ("仅中危待核项，建议确认后放行" if mid else "未发现明显缺陷")],
        [],
        ["规则", "命中数"],
    ]
    cnt = Counter(i[2] for i in issues)
    for r in ["R1", "R2", "R3", "R4", "R5", "R6", "R7", "R8"]:
        summary.append(["%s %s" % (r, RULE_NAME[r]), cnt.get(r, 0)])

    issues_sorted = sorted(issues, key=lambda x: SEV_ORDER.get(x[3], 9))
    sheet_issues = [["编号", "对象", "规则", "规则名", "严重度", "说明"]]
    for i in issues_sorted:
        sheet_issues.append([i[0], i[1], i[2], RULE_NAME[i[2]], i[3], i[4]])

    sheet_dev = [["维度", "目标", "实际", "提示"]]
    sheet_dev.extend(deviations)

    fixes = [["规则", "需补材料", "涉及对象数"]]
    per_rule = Counter(i[2] for i in issues)
    for r in ["R1", "R2", "R3", "R4", "R5", "R6", "R7", "R8"]:
        if per_rule.get(r):
            fixes.append([r + " " + RULE_NAME[r], RULE_FIX[r], per_rule[r]])

    return {
        "汇总": summary,
        "缺陷清单": sheet_issues,
        "偏离提示": sheet_dev,
        "待补材料": fixes,
    }


def write_xlsx(tables, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    first = True
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    for name, rows in tables.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 48)
    wb.save(out_path)


def ensure_dir(path):
    d = os.path.dirname(os.path.abspath(path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)


def write_csv(tables, out_path):
    base, _ = os.path.splitext(out_path)
    written = []
    for name, rows in tables.items():
        p = "%s_%s.csv" % (base, name)
        ensure_dir(p)
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        written.append(p)
    return written


def main():
    ap = argparse.ArgumentParser(description="受益对象名册体检")
    ap.add_argument("--input", required=True, help="输入 JSON 路径")
    ap.add_argument("--out", default="./体检报告.xlsx", help="输出 Excel 路径")
    ap.add_argument("--threshold", nargs=2, type=float, default=[0.6, 0.2],
                    metavar=("性别占比", "年龄超区间占比"), help="R7 提示阈值，默认 0.6 0.2")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print("输入文件不存在: %s" % args.input)
        sys.exit(2)

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    issues, deviations, total = run_rules(data, args.threshold[0], args.threshold[1])
    tables = build_tables(issues, deviations, total, data.get("project", {}))
    ensure_dir(args.out)

    try:
        write_xlsx(tables, args.out)
        print("已生成 Excel: %s" % args.out)
        print("名册 %d 人，命中 %d 条（高危 %d）" % (
            total, len(issues), sum(1 for i in issues if i[3] == "高")))
    except ImportError:
        paths = write_csv(tables, args.out)
        print("未安装 openpyxl，已降级输出 CSV：")
        for p in paths:
            print("  " + p)
    print("说明：命中项为待核项，不是违规结论，须人工确认后写入正式报告。")


if __name__ == "__main__":
    main()
